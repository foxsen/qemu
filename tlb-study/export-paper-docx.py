#!/usr/bin/env python3
"""Export the Chinese paper manuscript from Markdown to a Word document."""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HERE = Path(__file__).resolve().parent


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5):
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def configure_style(style, east_asia, latin, size, bold=False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def clean_inline(text):
    text = re.sub(r"!\[[^]]*\]\([^)]+\)", "", text)
    text = (text.replace("`", "").replace("**", "").replace("*", "")
            .replace("$", ""))
    text = text.replace("--", "—")
    return text.strip()


def equation_text(lines):
    text = " ".join(line.strip() for line in lines)
    replacements = {
        r"N_{refill}=N_{L1\ miss}-N_{victim\ hit}.":
            "N_refill = N_L1 miss − N_victim hit",
        r"p=p_b+(v-v_b),\quad 0\leq v-v_b<2^S. \tag{1}":
            "p = p_b + (v − v_b),   0 ≤ v − v_b < 2^S    (1)",
        r"S_{total}=\frac{1}{(1-f)+f/s}. \tag{2}":
            "S_total = 1 / [(1 − f) + f/s]    (2)",
    }
    return replacements.get(text, text.replace("\\quad", "   "))


def set_cell_text(cell, text, header=False):
    cell.text = clean_inline(text)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, size=9)
            run.bold = header


def add_table(document, lines):
    rows = [[part.strip() for part in line.strip().strip("|").split("|")]
            for line in lines]
    rows = [rows[0]] + rows[2:]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            set_cell_text(table.cell(row_index, column_index), value,
                          header=row_index == 0)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_body_paragraph(document, text, section_name):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(4)
    if section_name not in {"摘要", "Abstract", "参考文献"}:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    run = paragraph.add_run(clean_inline(text))
    set_run_font(run)
    return paragraph


def export(source, output):
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    configure_style(styles["Normal"], "宋体", "Times New Roman", 10.5)
    configure_style(styles["Title"], "黑体", "Times New Roman", 16, True)
    configure_style(styles["Heading 1"], "黑体", "Times New Roman", 13, True)
    configure_style(styles["Heading 2"], "黑体", "Times New Roman", 11, True)
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    configure_style(styles["Caption"], "宋体", "Times New Roman", 9, True)

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    title_count = 0
    section_name = ""
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(clean_inline(line[2:]))
            if title_count == 0:
                set_run_font(run, "黑体", "Times New Roman", 16)
            else:
                set_run_font(run, "黑体", "Times New Roman", 14)
            run.bold = True
            title_count += 1
            index += 1
            continue
        if line.startswith("## "):
            section_name = clean_inline(line[3:])
            document.add_heading(section_name, level=1)
            index += 1
            continue
        if line.startswith("### "):
            document.add_heading(clean_inline(line[4:]), level=2)
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            image_path = source.parent / image_match.group(2)
            if image_path.suffix.lower() == ".svg":
                image_path = image_path.with_name(
                    image_path.stem + "-300dpi.png")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Cm(16))
            index += 1
            continue
        if line.startswith("**图") or line.startswith("**表"):
            paragraph = document.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(clean_inline(line))
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, table_lines)
            continue
        if line == "$$":
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                equation_lines.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(equation_text(equation_lines))
            set_run_font(run, size=10.5)
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (not candidate or candidate.startswith("#") or
                    candidate.startswith("|") or candidate == "$$" or
                    candidate.startswith("![") or
                    candidate.startswith("**图") or
                    candidate.startswith("**表")):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = add_body_paragraph(document, " ".join(paragraph_lines),
                                       section_name)
        if paragraph.text.startswith("关键词：") or paragraph.text.startswith(
                "Keywords:"):
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.runs[0].bold = True

    document.core_properties.title = "全系统模拟器访存慢路径自适应优化"
    document.core_properties.subject = "高技术通讯投稿论文草稿"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path,
                        default=HERE / "PAPER_DRAFT_ZH.md")
    parser.add_argument("output", nargs="?", type=Path,
                        default=HERE / "PAPER_DRAFT_ZH.docx")
    args = parser.parse_args()
    export(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
